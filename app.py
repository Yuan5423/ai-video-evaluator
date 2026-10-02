import html
import io
import re
import textwrap
from datetime import datetime

import streamlit as st

from agent import evaluate_script

st.set_page_config(page_title="AI短视频口播评估助手", page_icon="🎬", layout="wide")

st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Noto+Sans+SC:wght@400;500;600;700&display=swap');
.stApp{background:#f5f6f8;color:#182230;font-family:'DM Sans','Noto Sans SC',sans-serif}.block-container{max-width:1160px;padding:46px 34px 72px}h1{font-size:2.2rem!important;letter-spacing:-.04em}h2{font-size:1.2rem!important;margin:28px 0 12px}.eyebrow{color:#635bff;font-size:.76rem;font-weight:700;letter-spacing:.14em;text-transform:uppercase;margin-bottom:12px}.hero-subtitle{color:#8993a3;font-size:1rem;margin-bottom:28px}.surface,.card,.candidate,.titlecard{background:#fff;border:1px solid #e8ebf1;border-radius:18px;box-shadow:0 8px 30px rgba(24,34,48,.045)}.surface{padding:24px}.card{padding:22px 24px;height:100%}.candidate{padding:25px;margin:12px 0}.titlecard{padding:17px 18px;min-height:105px}.label,.metric-label,.titlekind{color:#8791a0;font-size:.78rem;font-weight:600;margin:16px 0 6px}.copy{color:#394455;line-height:1.7}.metric{color:#172033;font-size:2rem;font-weight:700;letter-spacing:-.04em}.metric-note{color:#637083;font-size:.82rem;margin-top:4px}.score{background:linear-gradient(135deg,#f1efff,#f1f7ff);border-color:#dfddff}.score .metric,.rank,.cscore{color:#574cff}.hook{padding:25px 28px;border-left:4px solid #635bff;background:linear-gradient(110deg,#fff,#fafaff);border-radius:14px}.quote{color:#635bff;font-size:2rem;line-height:1}.hooktext{font-size:1.25rem;font-weight:600;line-height:1.65}.top{display:flex;justify-content:space-between;gap:16px}.rank{font-size:.76rem;font-weight:700;letter-spacing:.1em}.ctitle{font-size:1.15rem;font-weight:700;margin-top:5px}.cscore{font-size:1.35rem;font-weight:700}.flow{display:flex;gap:8px;align-items:center;flex-wrap:wrap}.step{background:#f2f4f8;color:#4b5666;border-radius:999px;padding:9px 14px;font-size:.83rem;font-weight:600}.arrow{color:#a0a8b5}.titletext{color:#263143;font-weight:600;line-height:1.5}.empty{padding:22px;background:#fafbfc;border:1px dashed #d9dee8;border-radius:14px;color:#7b8695}.history-item{padding:14px 0;border-bottom:1px solid #edf0f4}div[data-testid="stTextArea"] textarea{background:#fff;border:1px solid #e0e4eb;border-radius:14px;min-height:300px;padding:18px;font-size:1rem;line-height:1.7}div[data-testid="stTextArea"] textarea:focus{border-color:#9c96ff;box-shadow:0 0 0 3px rgba(99,91,255,.12)}div[data-testid="stButton"] button{border:0;border-radius:13px;background:linear-gradient(105deg,#5549e8,#3977f6);color:#fff;font-weight:700;min-height:46px;box-shadow:0 8px 18px rgba(73,79,220,.2)}
div[data-testid="stRadio"] label,div[data-testid="stRadio"] label span,div[data-testid="stRadio"] p{color:#263143!important;font-weight:600!important}
</style>""", unsafe_allow_html=True)

if "history" not in st.session_state:
    st.session_state.history = []
if "transcript_input" not in st.session_state:
    st.session_state.transcript_input = ""
if "active_report" not in st.session_state:
    st.session_state.active_report = None
if "active_image" not in st.session_state:
    st.session_state.active_image = None

def file_to_text(uploaded_file):
    if uploaded_file.name.lower().endswith(".docx"):
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("服务器尚未安装 Word 文件导入组件。") from exc
        document = Document(io.BytesIO(uploaded_file.getvalue()))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" ".join(cells))
        text = "\n".join(paragraphs)
        if not text:
            raise ValueError("该 Word 文档没有可读取的文字内容。")
        return text
    try:
        return uploaded_file.getvalue().decode("utf-8-sig").strip()
    except UnicodeDecodeError:
        return uploaded_file.getvalue().decode("gb18030").strip()

def format_highlighted_excerpt(text):
    """Safely render model-marked **key sentences** as bold text."""
    escaped = html.escape(text)
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)

def make_report_image(report):
    """Create a portable PNG summary for download and history preview."""
    from PIL import Image, ImageDraw, ImageFont
    font_candidates = [
        "C:/Windows/Fonts/msyh.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    font_path = next((path for path in font_candidates if __import__("os").path.exists(path)), None)
    regular = ImageFont.truetype(font_path, 34) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 27) if font_path else ImageFont.load_default()
    title = ImageFont.truetype(font_path, 46) if font_path else ImageFont.load_default()
    lines = ["AI口播剪辑价值评估报告", "", f"剪辑价值：{report.clip_value_level}", f"综合评分：{report.overall_score:.1f} / 10", f"有效可剪率：{report.effective_clip_rate:.1f}%", f"预计可剪条数：{report.suggested_video_count}", "", "总体结论：", report.overall_summary]
    lines.extend(["", "推荐开头（原文）：", report.best_opening.original_sentence, f"原文范围：{report.best_opening.source_range}", f"理由：{report.best_opening.reason}"])
    lines.extend(["", "可剪片段："])
    for index, candidate in enumerate(report.candidates, 1):
        plan = candidate.editing_plan
        ordered_lines = [f"第 {step} 句：{sentence}" for step, sentence in enumerate(plan.ordered_sentences, 1)]
        lines.extend([f"核心成片｜{candidate.clip_value_score:.1f} / 10", f"原文范围：{candidate.source_range}", f"核心观点：{candidate.core_viewpoint}", f"为什么值得剪：{candidate.reason}", f"独立完整度：{candidate.independence}", f"价值类型：{'、'.join(candidate.value_types)}", f"预计时长：{candidate.suggested_duration}", "剪辑方案（原文句子重排）：", *ordered_lines, f"重排说明：{plan.structure_reason}", ""])
    wrapped = []
    for line in lines:
        # Chinese transcripts often contain no spaces; force wrapping to avoid clipping text horizontally.
        wrapped.extend(textwrap.wrap(line, width=30, break_long_words=True, break_on_hyphens=False) or [""])
    image = Image.new("RGB", (1440, max(1000, 220 + len(wrapped) * 56)), "white")
    draw = ImageDraw.Draw(image)
    y = 60
    for index, line in enumerate(wrapped):
        draw.text((76, y), line, fill="#182230", font=title if index == 0 else (regular if index in (2, 3, 4) else small))
        y += 72 if index == 0 else 56
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()

def render_report(report):
    st.markdown('<div class="eyebrow" style="margin-top:36px">EVALUATION REPORT</div>', unsafe_allow_html=True)
    st.header("素材总评")
    st.markdown(f'<div class="surface"><div class="copy">{html.escape(report.overall_summary)}</div></div>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    c1, c2, c3 = st.columns([1.2, 1, 1])
    with c1:
        st.markdown(f'<div class="card score"><div class="metric-label">综合评分</div><div class="metric">{report.overall_score:.1f}<span style="font-size:1rem;color:#8d96a4"> / 10</span></div><div class="metric-note">{"达到剪辑标准" if report.overall_score >= 8 else "尚未达到剪辑标准"}</div></div>', unsafe_allow_html=True)
    with c2:
        st.markdown(f'<div class="card"><div class="metric-label">剪辑价值结论</div><div class="metric" style="font-size:1.35rem">{html.escape(report.clip_value_level)}</div><div class="metric-note">结合片段完整度与用户价值判断</div></div>', unsafe_allow_html=True)
    with c3:
        st.markdown(f'<div class="card"><div class="metric-label">有效可剪率</div><div class="metric">{report.effective_clip_rate:.1f}<span style="font-size:1rem;color:#8d96a4">%</span></div><div class="metric-note">预计可产出 {report.suggested_video_count} 条有效短视频</div></div>', unsafe_allow_html=True)

    st.header("五维评分")
    scores = report.overall_scores
    dimensions = [("开头吸引力", scores.hook_strength), ("观点清晰度", scores.viewpoint_clarity), ("信息密度", scores.information_density), ("情绪 / 冲突", scores.emotional_tension), ("传播潜力", scores.sharing_potential)]
    for col, (name, value) in zip(st.columns(5), dimensions):
        with col:
            st.markdown(f'<div class="card"><div class="metric-label">{name}</div><div class="metric" style="font-size:1.35rem">{value}<span style="font-size:.85rem;color:#a0a8b5"> / 2</span></div><div style="height:5px;background:#edf0f5;border-radius:4px;margin-top:12px"><div style="width:{value * 50}%;height:5px;background:linear-gradient(90deg,#635bff,#5c9bff);border-radius:4px"></div></div></div>', unsafe_allow_html=True)

    st.header("推荐开头")
    opening = report.best_opening
    st.markdown(f'<div class="hook"><div class="quote">“</div><div class="hooktext">{html.escape(opening.original_sentence)}</div><div class="label">原文范围</div><div class="copy">{html.escape(opening.source_range)}</div><div class="label">为什么适合前置</div><div class="copy">{html.escape(opening.reason)}</div></div>', unsafe_allow_html=True)

    st.header("值得剪的片段")
    if not report.candidates:
        st.markdown('<div class="empty">暂无明显值得单独剪出的片段。</div>', unsafe_allow_html=True)
    for index, candidate in enumerate(report.candidates, 1):
        plan = candidate.editing_plan
        ordered_steps = ''.join(f'<span class="step">第 {step} 句</span>{"<span class=\"arrow\">→</span>" if step < len(plan.ordered_sentences) else ""}' for step in range(1, len(plan.ordered_sentences) + 1))
        ordered_text = '<br>'.join(f'<b>第 {step} 句：</b>{html.escape(sentence)}' for step, sentence in enumerate(plan.ordered_sentences, 1))
        st.markdown(f'<div class="candidate"><div class="top"><div><div class="rank">核心成片 · {html.escape(candidate.source_range)}</div><div class="ctitle">{html.escape(candidate.core_viewpoint)}</div></div><div class="cscore">{candidate.clip_value_score:.1f}<span style="font-size:.8rem;color:#9aa3b0"> / 10</span></div></div><div class="label">使用的原文句子</div><div class="copy">{format_highlighted_excerpt(candidate.original_excerpt)}</div><div class="label">为什么值得剪</div><div class="copy">{html.escape(candidate.reason)}</div><div class="label">完整度 / 主要价值 / 预计时长</div><div class="copy">{candidate.independence} · {html.escape('、'.join(candidate.value_types))} · {html.escape(candidate.suggested_duration)}</div><div class="label">剪辑方案（原文句子重排）</div><div class="flow">{ordered_steps}</div><div class="copy">{ordered_text}<br><b>重排说明：</b>{html.escape(plan.structure_reason)}</div></div>', unsafe_allow_html=True)

    st.header("剪辑方案")
    if not report.candidates:
        st.markdown('<div class="empty">暂无可形成完整剪辑方案的独立片段。</div>', unsafe_allow_html=True)
    for index, candidate in enumerate(report.candidates, 1):
        plan = candidate.editing_plan
        ordered_steps = ''.join(f'<span class="step">第 {step} 句</span>{"<span class=\"arrow\">→</span>" if step < len(plan.ordered_sentences) else ""}' for step in range(1, len(plan.ordered_sentences) + 1))
        ordered_text = '<br>'.join(f'<b>第 {step} 句：</b>{html.escape(sentence)}' for step, sentence in enumerate(plan.ordered_sentences, 1))
        st.markdown(f'<div class="surface"><div class="rank">核心成片 · {html.escape(candidate.core_viewpoint)}</div><div class="flow">{ordered_steps}</div><div class="copy">{ordered_text}<br><b>重排说明：</b>{html.escape(plan.structure_reason)}</div></div>', unsafe_allow_html=True)

    image_bytes = st.session_state.get("active_image")
    if image_bytes:
        st.header("保存分析截图")
        st.download_button("下载分析截图（PNG）", data=image_bytes, file_name="ai-video-evaluation-report.png", mime="image/png", use_container_width=False)
        st.image(image_bytes, caption="当前分析报告截图预览", use_container_width=True)

st.markdown('<div class="eyebrow">AI CONTENT INTELLIGENCE</div>', unsafe_allow_html=True)
st.title("AI短视频口播评估助手")
st.markdown('<div class="hero-subtitle">从完整口播中识别真正值得剪辑的观点，让每一条内容都有清晰的传播理由。</div>', unsafe_allow_html=True)

with st.expander(f"历史评估记录（{len(st.session_state.history)}）", expanded=False):
    st.caption("记录仅保存在当前浏览器会话中；刷新浏览器、清除缓存或重新部署后会消失。")
    if not st.session_state.history:
        st.write("还没有历史记录。完成一次评估后会自动保存在这里。")
    else:
        history_indices = list(range(len(st.session_state.history) - 1, -1, -1))
        selected_index = st.selectbox("选择历史记录", history_indices, format_func=lambda i: f"{st.session_state.history[i]['created_at']} · {st.session_state.history[i]['score']:.1f}/10 · {st.session_state.history[i]['transcript'].replace(chr(10), ' ')[:36]}", key="selected_history_index")
        selected_item = st.session_state.history[selected_index]
        if st.button("查看选中的历史记录", key="view_selected_history", type="secondary"):
            st.session_state.transcript_input = selected_item["transcript"]
            st.session_state.active_report = selected_item["report"]
            st.session_state.active_image = selected_item.get("image")
            st.rerun()
        for index in history_indices:
            item = st.session_state.history[index]
            title = item["transcript"].replace("\n", " ").strip()[:42] or "未命名口播"
            left, right = st.columns([5, 1])
            with left:
                st.markdown(f"**{title}…**  \\n{item['created_at']} · {item['score']:.1f}/10")
            with right:
                if st.button("查看", key=f"history_{index}"):
                    st.session_state.transcript_input = item["transcript"]
                    st.session_state.active_report = item["report"]
                    st.session_state.active_image = item.get("image")
                    st.rerun()
        if st.button("清空本次历史", key="clear_history"):
            st.session_state.history = []
            st.session_state.active_report = None
            st.session_state.active_image = None
            st.rerun()

with st.container(border=True):
    st.markdown('<div class="label" style="margin-top:0">口播素材</div>', unsafe_allow_html=True)
    source_mode = st.radio("导入方式", ["导入 Word 文档", "粘贴文字"], horizontal=True, label_visibility="collapsed")
    if source_mode == "导入 Word 文档":
        uploaded = st.file_uploader("上传 Word 文档", type=["docx"], help="仅支持未加密的 .docx 文件。旧版 .doc 请先在 Word 中另存为 .docx。")
        st.caption("上传后会自动提取 Word 正文和表格文字，无需再粘贴内容。")
        if uploaded is not None:
            try:
                uploaded_text = file_to_text(uploaded)
            except ValueError as exc:
                st.error(str(exc))
            except Exception:
                st.error("Word 文档读取失败。请确认文件未加密、未损坏，并且是 .docx 格式；旧版 .doc 文件请先另存为 .docx。")
            else:
                fingerprint = f"{uploaded.name}:{uploaded.size}"
                if st.session_state.get("uploaded_fingerprint") != fingerprint:
                    st.session_state.transcript_input = uploaded_text
                    st.session_state.uploaded_fingerprint = fingerprint
                st.success(f"已读取《{uploaded.name}》的文字内容，点击下方按钮即可开始评估。")
    else:
        st.text_area("口播转写稿", key="transcript_input", height=330, label_visibility="collapsed", placeholder="请粘贴完整的口播转写稿，不需要自己整理，我们会自动寻找值得剪辑的内容。")
        st.caption("我们会阅读全文，并从任何位置寻找最强开头。")
    run = st.button("开始评估  →", type="primary", use_container_width=True)

if run:
    transcript = st.session_state.transcript_input
    if not transcript.strip():
        st.warning("请先输入或导入口播转写稿。")
    else:
        with st.spinner("正在完整阅读并寻找高价值片段…"):
            try:
                report = evaluate_script(transcript)
            except Exception as exc:
                st.error(f"评估失败（{type(exc).__name__}）：{exc}")
                st.caption("配置检测、网络、模型限流或模型返回格式都可能导致失败；上方信息可用于准确定位。")
            else:
                st.session_state.active_report = report
                report_image = make_report_image(report)
                st.session_state.active_image = report_image
                st.session_state.history.append({
                    "transcript": transcript,
                    "report": report,
                    "score": report.overall_score,
                    "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
                    "image": report_image,
                })
                st.success("评估完成，结果已保存到本次历史记录。")

if st.session_state.active_report is not None:
    render_report(st.session_state.active_report)

